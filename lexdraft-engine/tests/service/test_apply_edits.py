"""Edited text written back into the engine .docx keeps the court layout."""

import base64
import io

import pytest

pytest.importorskip("fastapi", reason="service extras not installed (pip install -e .[service])")

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from src.service.app import app  # noqa: E402

client = TestClient(app)


def _blank():
    r = client.post("/v1/blank", json={"doc_type": "urgent_petition"}).json()
    return r["docx_base64"], [b["text"] for b in r["blocks"]]


def test_filled_blanks_land_in_the_original_paragraphs():
    docx, blocks = _blank()
    edited = [b.replace("____________", "Anand", 1) if "Petitioner" in b else b for b in blocks]
    r = client.post("/v1/apply-edits", json={"docx_base64": docx, "paragraphs": edited}).json()
    assert r["matched"] is True
    out = Document(io.BytesIO(base64.b64decode(r["docx_base64"])))
    before = Document(io.BytesIO(base64.b64decode(docx)))
    assert any("Anand" in p.text and "Petitioner" in p.text for p in out.paragraphs)
    # layout kept: same paragraph count, same alignment, same page setup
    assert len(out.paragraphs) == len(before.paragraphs)
    assert [p.alignment for p in out.paragraphs] == [p.alignment for p in before.paragraphs]
    assert out.sections[0].left_margin == before.sections[0].left_margin


def test_added_paragraph_falls_back():
    docx, blocks = _blank()
    r = client.post("/v1/apply-edits", json={"docx_base64": docx, "paragraphs": blocks + ["A new paragraph."]})
    assert r.json() == {"matched": False}


def test_unchanged_text_round_trips():
    docx, blocks = _blank()
    r = client.post("/v1/apply-edits", json={"docx_base64": docx, "paragraphs": blocks}).json()
    assert r["matched"] is True
