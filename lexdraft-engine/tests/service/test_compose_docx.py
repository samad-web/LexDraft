"""Plain drafted text -> court-format .docx (house style, recognised shapes)."""

import base64
import io

import pytest

pytest.importorskip("fastapi", reason="service extras not installed (pip install -e .[service])")

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from src.service.app import app  # noqa: E402
from src.service.compose import classify  # noqa: E402

client = TestClient(app)
LINES = [
    "IN THE FAMILY COURT AT CHENNAI",
    "H.M.O.P. No. [_____] of 2026",
    "Sivaranjini, W/o. Raghuram, residing at [address].",
    "... PETITIONER",
    "VERSUS",
    "Raghuram, S/o. [father's name], residing at [address].",
    "... RESPONDENT",
    "MOST RESPECTFULLY SHOWETH:",
    "1. The marriage was solemnised on 2004-03-24 at [place of marriage] according to Hindu rites.",
    "PRAYER",
    "It is therefore prayed that this Hon'ble Court may be pleased to dissolve the marriage.",
    "Counsel for the Petitioner",
]


def test_shapes_are_recognised():
    styles = [b["style"] for b in classify(LINES)]
    assert styles == ["title", "center", "left", "right", "center", "left", "right",
                      "heading", "para", "heading", "para", "signature"]


def test_endpoint_returns_court_format_docx():
    r = client.post("/v1/compose-docx", json={"paragraphs": LINES})
    assert r.status_code == 200, r.text
    body = r.json()
    doc = Document(io.BytesIO(base64.b64decode(body["docx_base64"])))
    sec = doc.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0 and round(sec.left_margin.cm, 2) == 3.81
    by_text = {p.text: p for p in doc.paragraphs if p.text}
    assert by_text["IN THE FAMILY COURT AT CHENNAI"].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert by_text["... PETITIONER"].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    # dates tidied to the Indian court convention
    assert any("24.03.2004" in p.text for p in doc.paragraphs)
    assert len(body["blocks"]) == len(LINES)


def test_rejects_bad_input():
    assert client.post("/v1/compose-docx", json={"paragraphs": []}).status_code == 422
