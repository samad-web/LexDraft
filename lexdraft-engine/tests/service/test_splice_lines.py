"""splice_lines: write a composed body back into an engine .docx, keeping the
cause title above and the prayer/verification below exactly as rendered."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.service.apply_edits import splice_lines


def _doc(paras):
    d = Document()
    for text, align, bold in paras:
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if align:
            p.alignment = align
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _texts(raw):
    return [p.text for p in Document(io.BytesIO(raw)).paragraphs if p.text.strip()]


SRC = _doc([
    ("IN THE COURT OF THE JUDICIAL MAGISTRATE", WD_ALIGN_PARAGRAPH.CENTER, True),  # line 0
    ("1. The petitioner was arrested on 10.10.2026.", WD_ALIGN_PARAGRAPH.JUSTIFY, False),  # 1
    ("2. He is innocent.", WD_ALIGN_PARAGRAPH.JUSTIFY, False),  # 2
    ("PRAYER", WD_ALIGN_PARAGRAPH.CENTER, True),  # 3
    ("Place: Chennai", None, False),  # 4
])


def test_more_paragraphs_keeps_head_and_tail():
    body = ["1. Facts A.", "2. Facts B.", "3. Ground C.", "4. Ground D."]
    out = splice_lines(SRC, 1, 2, body)
    assert out is not None
    assert _texts(out) == ["IN THE COURT OF THE JUDICIAL MAGISTRATE", *body, "PRAYER", "Place: Chennai"]
    paras = [p for p in Document(io.BytesIO(out)).paragraphs if p.text.strip()]
    # Cloned body paragraphs inherit the body style, not the heading's.
    assert paras[4].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert not paras[4].runs[0].bold
    assert paras[0].runs[0].bold and paras[5].runs[0].bold


def test_fewer_paragraphs_removes_surplus():
    out = splice_lines(SRC, 1, 2, ["1. Everything in one paragraph."])
    assert _texts(out) == [
        "IN THE COURT OF THE JUDICIAL MAGISTRATE", "1. Everything in one paragraph.", "PRAYER", "Place: Chennai",
    ]


def test_blank_spacers_inside_the_range_are_removed():
    d = Document()
    for t in ["HEAD", "1. A.", "", "", "2. B.", "", "PRAYER"]:
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    out = splice_lines(buf.getvalue(), 1, 2, ["1. New A.", "2. New B.", "3. New C."])
    paras = [p.text for p in Document(io.BytesIO(out)).paragraphs]
    # The spacers between the old body lines are gone; the one before PRAYER stays.
    assert paras[-6:] == ["HEAD", "1. New A.", "2. New B.", "3. New C.", "", "PRAYER"]


def test_bad_ranges_return_none():
    assert splice_lines(SRC, 3, 1, ["x"]) is None
    assert splice_lines(SRC, 1, 9, ["x"]) is None
    assert splice_lines(SRC, 1, 2, ["   "]) is None


def test_range_cutting_a_soft_broken_paragraph_returns_none():
    d = Document()
    d.add_paragraph("HEAD")
    p = d.add_paragraph()
    p.add_run("line one")
    p.add_run().add_break()
    p.add_run("line two")
    buf = io.BytesIO()
    d.save(buf)
    # Lines: 0 HEAD, 1 line one, 2 line two — replacing only line 1 would split it.
    assert splice_lines(buf.getvalue(), 1, 1, ["x"]) is None
    assert splice_lines(buf.getvalue(), 1, 2, ["x"]) is not None


def test_body_never_inherits_a_bold_sub_heading():
    # The s.138 complaint skeleton opens its body with BOLD sub-headings
    # ("1. Description of the Complainant"); the composed pleading spliced over
    # them must come out as plain justified body text, not bold.
    src = _doc([
        ("IN THE COURT OF THE JUDICIAL MAGISTRATE", WD_ALIGN_PARAGRAPH.CENTER, True),
        ("1. Description of the Complainant", None, True),
        ("Abdul Rahman, having his office at Chennai.", WD_ALIGN_PARAGRAPH.JUSTIFY, False),
        ("2. Description of the Accused", None, True),
        ("3. The Complainant humbly submits that the cheque was dishonoured.", WD_ALIGN_PARAGRAPH.JUSTIFY, False),
        ("PRAYER", WD_ALIGN_PARAGRAPH.CENTER, True),
    ])
    body = ["1. The Complainant, Mr. Abdul Rahman, is the payee.", "2. The Accused is the drawer.",
            "FACTS OF THE CASE", "3. On 10 January 2026 the Accused borrowed Rs. 5,00,000/-."]
    out = splice_lines(src, 1, 4, body)
    paras = [p for p in Document(io.BytesIO(out)).paragraphs if p.text.strip()]
    by_text = {p.text: p for p in paras}
    for t in (body[0], body[1], body[3]):
        p = by_text[t]
        assert not any(r.bold for r in p.runs), t
        assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, t
    assert any(r.bold for r in by_text["FACTS OF THE CASE"].runs)
    assert [p.text for p in paras][0].startswith("IN THE COURT")
    assert [p.text for p in paras][-1] == "PRAYER"
    assert any(r.bold for r in paras[0].runs) and any(r.bold for r in paras[-1].runs)
