"""Imported formats: manifest -> generic adapter -> template -> HTTP draft.

Pins the guarantees the importer relies on:
  - an imported format drafts through the SAME service boundary as a built-in
    (422 on missing required input, generated docx otherwise);
  - a pending format carries the "not yet approved" marker, an approved one doesn't;
  - an imported entry can never shadow a built-in doc type (slug or title);
  - a spec with an undeclared / malformed placeholder is rejected before it is built;
  - the LAWE catalog only points at doc types the engine actually has.
"""

import base64
import copy
import io
from pathlib import Path

import pytest

from docx import Document

from src.components import render
from src.imports import builder, lawe
from src.imports import manifest as mf
from src.service import adapters
from src.service.adapters import imported

SPEC = {
    "slug": "test_hc_summons",
    "title": "Test HC Original Side Summons",
    "court_issued": True,
    "fields": [
        {"key": "case_ref", "kind": "string", "required": True, "label": "Suit number",
         "help": "", "sample": "C.S. No. 12 of 2026"},
        {"key": "plaintiff_name", "kind": "string", "required": True, "label": "Plaintiff",
         "help": "", "sample": "Anand"},
        {"key": "defendant_name", "kind": "string", "required": True, "label": "Defendant",
         "help": "", "sample": "Balu"},
        {"key": "plaint_date", "kind": "date", "required": False, "label": "Plaint presented on",
         "help": "", "sample": "2026-03-02"},
        {"key": "leave_granted", "kind": "bool", "required": False, "label": "Leave granted?",
         "help": "", "sample": "true"},
        {"key": "documents", "kind": "string_list", "required": False, "label": "Documents",
         "help": "", "sample": ["Plaint copy", "Document list"]},
    ],
    "blocks": [
        {"style": "title", "text": "IN THE HIGH COURT OF JUDICATURE AT MADRAS", "bold": True, "condition": ""},
        {"style": "center", "text": "{{case_ref}}", "bold": True, "condition": ""},
        {"style": "left", "text": "{{ plaintiff_name }} ... Plaintiff", "bold": False, "condition": ""},
        {"style": "center", "text": "and", "bold": False, "condition": ""},
        {"style": "left", "text": "{{ defendant_name }} ... Defendant", "bold": False, "condition": ""},
        {"style": "para", "text": "Whereas the plaintiff presented a plaint on {{ plaint_date }}.",
         "bold": False, "condition": ""},
        {"style": "para", "text": "Leave to institute this suit was granted.", "bold": False,
         "condition": "leave_granted"},
        {"style": "heading", "text": "DOCUMENTS", "bold": True, "condition": ""},
        {"style": "list", "text": "documents", "bold": False, "condition": ""},
        {"style": "signature", "text": "Assistant Registrar", "bold": False, "condition": ""},
    ],
}


def _entry(status="pending_review", slug="test_hc_summons", title=None):
    return {
        "title": title or SPEC["title"], "status": status, "version": 1,
        "template": f"{slug}.docx", "court_issued": True,
        "fields": copy.deepcopy(SPEC["fields"]),
        "source": {"kind": "lawe", "ref": "https://example.test/x.pdf"},
    }


@pytest.fixture
def imported_dir(tmp_path, monkeypatch):
    """Point the imported-template directory at a temp dir holding the test template."""
    monkeypatch.setattr(render, "_IMPORTED_DIR", tmp_path)
    monkeypatch.setattr(mf, "IMPORTED_DIR", tmp_path)
    builder.build_docx(SPEC, tmp_path / "test_hc_summons.docx")
    return tmp_path


@pytest.fixture
def registered(imported_dir, monkeypatch):
    """Register test entries into the live registry; remove them afterwards.
    Serves pending entries (as local dev does) unless a test unsets it."""
    monkeypatch.setenv("LEXDRAFT_IMPORTS_SERVE_PENDING", "true")
    added = []

    def _register(manifest):
        done = imported.register_all(manifest)
        added.extend(done)
        return done

    yield _register
    for slug in added:
        adapters._ADAPTERS.pop(slug, None)
        imported.IMPORTED_DOC_TYPES.discard(slug)
        imported._SAMPLES.pop(slug, None)


def _text(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


# -- spec validation ---------------------------------------------------------

def test_valid_spec_has_no_problems():
    assert builder.check_spec(SPEC) == []


@pytest.mark.parametrize("mutate, expect", [
    (lambda s: s["blocks"].append({"style": "para", "text": "{{ ghost }}", "bold": False,
                                   "condition": ""}), "no declared field"),
    (lambda s: s["blocks"].append({"style": "para", "text": "{% if x %}x{% endif %}",
                                   "bold": False, "condition": ""}), "control tags"),
    (lambda s: s["blocks"].append({"style": "para", "text": "{{ broken", "bold": False,
                                   "condition": ""}), "malformed placeholder"),
    (lambda s: s["blocks"].__setitem__(6, dict(s["blocks"][6], condition="case_ref")),
     "must name a bool field"),
    (lambda s: s["fields"].append({"key": "unused", "kind": "string", "required": False,
                                   "label": "", "help": "", "sample": ""}), "never placed"),
    (lambda s: s["fields"].append({"key": "review_banner", "kind": "string", "required": False,
                                   "label": "", "help": "", "sample": ""}), "reserved"),
    (lambda s: s["blocks"].__setitem__(8, dict(s["blocks"][8], text="case_ref")),
     "string_list field"),
])
def test_bad_spec_is_rejected(mutate, expect):
    spec = copy.deepcopy(SPEC)
    mutate(spec)
    problems = builder.check_spec(spec)
    assert any(expect in p for p in problems), problems


def test_smoke_render_is_clean(imported_dir):
    assert builder.smoke_render("test_hc_summons", _entry(), imported_dir / "test_hc_summons.docx") == []


def test_template_page_setup_is_a4_with_binding_margin(imported_dir):
    section = Document(str(imported_dir / "test_hc_summons.docx")).sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 2) == 3.81


# -- adapter behaviour ---------------------------------------------------------

def test_pending_format_carries_marker_and_blanks(imported_dir):
    adapter = imported.adapter_for("test_hc_summons", _entry())
    kwargs = adapter.coerce({"case_ref": "C.S. 1/2026", "plaintiff_name": "Anand",
                             "defendant_name": "Balu"})
    ctx = adapter.build(**kwargs).context
    assert "not yet approved" in ctx["unverified_marker"]
    assert ctx["plaint_date"] == imported.BLANK
    assert ctx["documents"] == [] and ctx["leave_granted"] is False
    assert ctx["court_issues"] is True


def test_approved_format_has_no_marker(imported_dir):
    ctx = imported.adapter_for("test_hc_summons", _entry(status="approved")).build(
        case_ref="x", plaintiff_name="a", defendant_name="b").context
    assert ctx["unverified_marker"] == ""


def test_missing_required_is_a_422_not_a_refusal(imported_dir):
    adapter = imported.adapter_for("test_hc_summons", _entry())
    with pytest.raises(adapters.AdapterValidationError) as exc:
        adapter.coerce({"case_ref": "x"})
    assert {i.field for i in exc.value.issues} == {"plaintiff_name", "defendant_name"}


def test_register_never_shadows_a_builtin(registered):
    before = {a.doc_type: a for a in adapters.all_adapters()}
    manifest = {"doc_types": {
        "urgent_petition": _entry(slug="test_hc_summons"),                  # built-in slug
        "test_title_clash": _entry(title="Urgent Petition (s.151 CPC)"),    # built-in title
        "test_hc_summons": _entry(),                                        # fine
        "test_no_template": dict(_entry(), template="missing.docx"),        # no file
        "Bad Slug": _entry(),                                               # invalid
    }}
    assert registered(manifest) == ["test_hc_summons"]
    assert adapters.get_adapter("urgent_petition") is before["urgent_petition"]


def test_pending_formats_are_hidden_by_default(registered, monkeypatch):
    # Production default: only approved formats are served.
    monkeypatch.delenv("LEXDRAFT_IMPORTS_SERVE_PENDING")
    assert registered({"doc_types": {"test_hc_summons": _entry()}}) == []
    assert registered({"doc_types": {"test_hc_summons": _entry(status="approved")}}) == ["test_hc_summons"]


def test_template_resolution(imported_dir):
    assert render.template_path("imported/test_hc_summons.docx") == imported_dir / "test_hc_summons.docx"
    assert render.template_path("urgent_petition.docx").parent.name == "shared"


# -- through HTTP --------------------------------------------------------------

def test_draft_endpoint_renders_imported_format(registered):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient
    from src.service.app import app

    registered({"doc_types": {"test_hc_summons": _entry()}})
    client = TestClient(app)
    described = next(d for d in client.get("/v1/doc-types").json()["doc_types"]
                     if d["doc_type"] == "test_hc_summons")
    assert {f["key"] for f in described["fields"]} == {f["key"] for f in SPEC["fields"]}

    r = client.post("/v1/draft/test_hc_summons", json={
        "matter": {"matter_id": "M-IMP-1", "version": 1},
        "fields": {"case_ref": "C.S. No. 12 of 2026", "plaintiff_name": "Anand",
                   "defendant_name": "Balu", "leave_granted": True,
                   "documents": ["Plaint copy"]},
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["outcome"] == "generated"
    assert body["template_id"] == "imported/test_hc_summons.docx"
    text = _text(base64.b64decode(body["docx_base64"]))
    assert "Anand ... Plaintiff" in text and "C.S. No. 12 of 2026" in text
    assert "Leave to institute this suit was granted." in text
    assert "Plaint copy" in text and "not yet approved" in text
    assert "{{" not in text and "{%" not in text

    missing = client.post("/v1/draft/test_hc_summons", json={
        "matter": {"matter_id": "M-IMP-1", "version": 1}, "fields": {"case_ref": "x"}})
    assert missing.status_code == 422


# -- LAWE ----------------------------------------------------------------------

def test_lawe_listing_parser():
    js = ('download[ndx] = "typeofdoc~foldername~filename"\n'
          'download[++ndx] = "pdf~High_Court~Caveat_Petition.pdf"\n'
          'download[++ndx] = "pdf~Others~Bail_Bond.pdf"\n')
    assert lawe.parse_listing(js) == [("High_Court", "Caveat_Petition.pdf"), ("Others", "Bail_Bond.pdf")]
    assert lawe.title_of("Appellete_Side_-_Copy_Application.pdf") == "Appellete Side - Copy Application"


def test_lawe_catalog_points_at_real_doc_types():
    registered = {a.doc_type for a in adapters.all_adapters()}
    imported_slugs = set(mf.load_manifest()["doc_types"])   # pending ones aren't registered
    forms = lawe.load_catalog()["forms"]
    assert forms, "catalog is empty"
    for fid, c in forms.items():
        assert c["status"] in ("covered", "imported", "not_a_format"), fid
        if c["status"] == "covered":
            assert c["doc_type"] in registered, f"{fid} -> {c['doc_type']} is not a doc type"
        elif c["status"] == "imported":
            assert c["doc_type"] in imported_slugs, f"{fid} -> {c['doc_type']} not in the manifest"


def test_committed_manifest_entries_are_valid():
    for slug, entry in mf.load_manifest()["doc_types"].items():
        assert mf.validate_entry(slug, entry) == [], slug
        assert (mf.IMPORTED_DIR / entry["template"]).exists(), slug
