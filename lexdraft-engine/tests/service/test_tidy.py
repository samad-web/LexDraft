"""tidy_docx: no leading blanks, no runs of blank paragraphs, Indian body dates."""

from docx import Document

import pytest

from src.service.tidy import court_heading, tidy_docx


@pytest.mark.parametrize("raw, want", [
    ("IN THE COURT OF Judicial Magistrate No. II, Egmore, Chennai (DISTRICT Chennai), ",
     "IN THE COURT OF JUDICIAL MAGISTRATE NO. II, EGMORE, CHENNAI"),
    ("IN THE COURT OF METROPOLITAN MAGISTRATE (DISTRICT Salem), Salem",
     "IN THE COURT OF METROPOLITAN MAGISTRATE (DISTRICT SALEM)"),
    ("IN THE COURT OF THE CHIEF METROPOLITAN MAGISTRATE, Saidapet (DISTRICT Chennai), Chennai",
     "IN THE COURT OF THE CHIEF METROPOLITAN MAGISTRATE, SAIDAPET (DISTRICT CHENNAI)"),
    ("IN THE COURT OF THE DISTRICT MUNSIF, TAMBARAM", "IN THE COURT OF THE DISTRICT MUNSIF, TAMBARAM"),
])
def test_court_heading(raw, want):
    assert court_heading(raw) == want


def test_collapses_blanks_and_formats_dates(tmp_path):
    d = Document()
    for t in ["", "", "TITLE", "", "", "", "Taken into custody on 2026-10-10.", "", "Date: 2026-10-26"]:
        d.add_paragraph(t)
    table = d.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "2026-10-10"
    path = tmp_path / "x.docx"
    d.save(path)

    tidy_docx(path)

    doc = Document(path)
    assert [p.text for p in doc.paragraphs] == [
        "TITLE", "", "Taken into custody on 10.10.2026.", "", "Date: 26.10.2026",
    ]
    assert doc.tables[0].rows[0].cells[0].text == "2026-10-10"  # tables untouched


@pytest.mark.parametrize("digits, want", [
    ("500", "500"), ("5000", "5,000"), ("500000", "5,00,000"),
    ("1000000", "10,00,000"), ("125000000", "12,50,00,000"),
])
def test_indian_grouping(digits, want):
    from src.service.tidy import indian_grouping
    assert indian_grouping(digits) == want


def test_amounts_dates_and_numbering_gap(tmp_path):
    # The s.138 complaint as served: an optional para 6 rendered empty, raw
    # figures and dd-mm-yyyy dates in the body.
    d = Document()
    for t in ["5. The cheque was presented.", "7. Owing to the dishonour, Rs. 500000/- is due.",
              "8. Twice the amount, being Rs. 1000000/-, already 5,00,000 grouped.",
              "PRAYER", "1. Dated 10-03-2026 — Cheque", "2. Return memo", "4. Notice"]:
        d.add_paragraph(t)
    path = tmp_path / "x.docx"
    d.save(path)
    tidy_docx(path)
    got = [p.text for p in Document(path).paragraphs]
    assert got[0] == "5. The cheque was presented."
    assert got[1] == "6. Owing to the dishonour, Rs. 5,00,000/- is due."
    assert got[2] == "7. Twice the amount, being Rs. 10,00,000/-, already 5,00,000 grouped."
    assert got[4] == "1. Dated 10.03.2026 — Cheque"   # a list restarts at 1, untouched
    assert got[6] == "3. Notice"
