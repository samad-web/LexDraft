"""Blind drafting benchmark — score drafts from several tools on the same cases.

Every tool gets the SAME brief from cases.json. Save each tool's output as

    <outputs>/<tool>/<case_id>.docx | .pdf | .txt

(e.g. outputs/lexdraft/s138_notice.docx, outputs/draftbotpro/s138_notice.pdf). If a
tool refused or warned instead of drafting, save that message as <case_id>.txt.

Scores per draft (0-100, weights in WEIGHTS):
  averments     share of the document's MANDATORY averments (verify/averments/)
                actually pleaded. Needs --grade (an LLM judge that sees only the
                text — never which tool wrote it).
  facts         share of the case's key facts reproduced correctly.
  old_law       no repealed IPC / CrPC / Evidence Act references.
  placeholders  no unfilled blanks — every fact needed was given in the brief.
  figures       no amounts that appear nowhere in the brief (invented numbers).
  diction       no high-severity diction slips (src/verify/diction.py).
Trap cases (expect = flag) are scored separately: did the tool catch the planted
defect instead of drafting as if it were fine? (needs --grade).

Dates and case-law citations not in the brief are LISTED for a human to check,
not penalised — derived dates (cause of action, limitation) are legitimate, and
citations can only be verified against a judgment database.

Usage:
    python tools/benchmark/score.py --outputs bench_out
    python tools/benchmark/score.py --outputs bench_out --grade      # needs ANTHROPIC_API_KEY
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import urllib.request
from pathlib import Path
from typing import Dict, List, Optional

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.verify import averments as _averments  # noqa: E402
from src.verify import diction as _diction  # noqa: E402

CASES_FILE = Path(__file__).with_name("cases.json")
WEIGHTS = {"averments": 40, "facts": 25, "old_law": 10, "placeholders": 10, "figures": 10, "diction": 5}
EXTS = (".docx", ".pdf", ".txt", ".md")

_OLD_LAW = re.compile(
    r"\bI\.?\s?P\.?\s?C\b|Indian Penal Code|\bCr\.?\s?P\.?\s?C\b|Code of Criminal Procedure|Indian Evidence Act",
    re.IGNORECASE,
)
_PLACEHOLDER = re.compile(
    r"_{3,}|\[[^\]\n]{0,60}(?:insert|name|date|amount|address|details|fill|number|xx)[^\]\n]{0,60}\]|\bX{3,}\b",
    re.IGNORECASE,
)
_AMOUNT = re.compile(r"(?:₹|\bRs\.?|\bINR)\s*([\d,]+(?:\.\d+)?)", re.IGNORECASE)
_DATE = re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4})\b")
_CITATION = re.compile(
    r"\(\d{4}\)\s*\d+\s*SCC\s*\d+|\d{4}\s*SCC\s*OnLine\s*[A-Za-z]+\s*\d+|AIR\s*\d{4}\s*[A-Z][A-Za-z]*\s*\d+"
)


# ---- text extraction -------------------------------------------------------

def load_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    return path.read_text(encoding="utf-8", errors="replace")


def _norm(s: str) -> str:
    return re.sub(r"[\s.,\-/()'\"’‘]+", "", s.lower())


def _digits(s: str) -> str:
    return re.sub(r"\D", "", s.split(".")[0])


# ---- deterministic checks --------------------------------------------------

def fact_score(text: str, key_facts: List) -> Dict:
    hay = _norm(text)
    missing = []
    for fact in key_facts:
        options = fact if isinstance(fact, list) else [fact]
        if not any(_norm(o) in hay for o in options):
            missing.append(options[0])
    found = len(key_facts) - len(missing)
    return {"score": found / len(key_facts) if key_facts else 1.0, "missing": missing}


def unsupported_amounts(text: str, brief: str) -> List[str]:
    allowed = {_digits(m) for m in _AMOUNT.findall(brief)}
    out = []
    for m in _AMOUNT.finditer(text):
        d = _digits(m.group(1))
        if d and d not in allowed and m.group(0).strip() not in out:
            out.append(m.group(0).strip())
    return out


def unlisted_dates(text: str, brief: str) -> List[str]:
    key = lambda m: (int(m[0]), int(m[1]), int(m[2]))  # noqa: E731
    allowed = {key(m) for m in _DATE.findall(brief)}
    seen, out = set(), []
    for m in _DATE.finditer(text):
        k = key(m.groups())
        if k not in allowed and k not in seen:
            seen.add(k)
            out.append(m.group(0))
    return out


def deterministic(text: str, case: Dict) -> Dict:
    facts = fact_score(text, case["key_facts"])
    old_law = sorted({m.group(0) for m in _OLD_LAW.finditer(text)})
    placeholders = _PLACEHOLDER.findall(text)
    amounts = unsupported_amounts(text, case["brief"])
    diction_high = [f.message for f in _diction.lint(text) if f.severity == "high"]
    return {
        "facts": facts,
        "old_law": old_law,
        "placeholders": len(placeholders),
        "unsupported_amounts": amounts,
        "unlisted_dates": unlisted_dates(text, case["brief"]),
        "citations": sorted(set(_CITATION.findall(text))),
        "diction_high": diction_high,
        "chars": len(text),
    }


# ---- LLM judge (optional) --------------------------------------------------

_JUDGE_SYSTEM = (
    "You are a strict senior Indian advocate grading a draft legal document. You see only the "
    "document text. For each numbered requirement, decide if the document actually pleads it "
    "(in substance, not just a heading). A requirement that does not apply on these facts "
    "(e.g. a company-only averment when no company is involved) counts as present. "
    "Return ONLY JSON: {\"results\":[{\"i\":0,\"present\":true,\"evidence\":\"short quote or why missing\"}],"
    "\"defect_flagged\":null}"
)


def _claude(system: str, user: str) -> str:
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        raise SystemExit("--grade needs ANTHROPIC_API_KEY in the environment")
    body = json.dumps({
        "model": os.environ.get("BENCH_MODEL", "claude-opus-5-5"),
        "max_tokens": 4000,
        "system": system,
        "messages": [{"role": "user", "content": user}],
    }).encode()
    req = urllib.request.Request(
        "https://api.anthropic.com/v1/messages", data=body, method="POST",
        headers={"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=300) as r:
        data = json.loads(r.read())
    return "".join(b.get("text", "") for b in data.get("content", []))


def _loose_json(raw: str) -> Optional[Dict]:
    raw = re.sub(r"```(?:json)?", "", raw)
    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        return json.loads(raw[start:end + 1])
    except json.JSONDecodeError:
        return None


def judge(text: str, case: Dict) -> Dict:
    mandatory = _averments.get(case["averments"]).mandatory()
    reqs = "\n".join(f"[{i}] {a.averment}" for i, a in enumerate(mandatory))
    trap = ""
    if case["expect"] == "flag":
        trap = (
            "\n\nALSO: set \"defect_flagged\" to true ONLY if the text explicitly identifies this legal "
            f"defect (refusing, warning, or advising another remedy), else false:\n{case['trap_issue']}"
        )
    user = f"FACTS GIVEN TO THE DRAFTER:\n{case['brief']}\n\nREQUIREMENTS:\n{reqs}{trap}\n\nDOCUMENT:\n{text[:120_000]}"
    parsed = _loose_json(_claude(_JUDGE_SYSTEM, user)) or {}
    results = {r.get("i"): r for r in parsed.get("results", []) if isinstance(r, dict)}
    missing = [mandatory[i].averment for i in range(len(mandatory)) if not results.get(i, {}).get("present")]
    return {
        "score": (len(mandatory) - len(missing)) / len(mandatory) if mandatory else 1.0,
        "missing": missing,
        "defect_flagged": parsed.get("defect_flagged"),
    }


# ---- scoring ---------------------------------------------------------------

def overall(det: Dict, avr: Optional[Dict]) -> float:
    parts = {
        "facts": det["facts"]["score"],
        "old_law": 0.0 if det["old_law"] else 1.0,
        "placeholders": max(0.0, 1 - 0.25 * det["placeholders"]),
        "figures": max(0.0, 1 - 0.34 * len(det["unsupported_amounts"])),
        "diction": max(0.0, 1 - 0.5 * len(det["diction_high"])),
    }
    if avr is not None:
        parts["averments"] = avr["score"]
    total_w = sum(WEIGHTS[k] for k in parts)
    return round(100 * sum(WEIGHTS[k] * v for k, v in parts.items()) / total_w, 1)


def find_output(tool_dir: Path, case_id: str) -> Optional[Path]:
    for ext in EXTS:
        p = tool_dir / f"{case_id}{ext}"
        if p.exists():
            return p
    return None


def render_report(cases: List[Dict], tools: List[str], res: Dict, graded: bool) -> str:
    lines = ["# Drafting benchmark", ""]
    lines.append("Scores are 0-100 (" + ", ".join(f"{k} {v}" for k, v in WEIGHTS.items())
                 + ")." + ("" if graded else " Averments NOT graded (run with --grade); weights re-normalised."))
    lines += ["", "## Summary", "", "| Tool | Drafts scored | Mean score | Traps caught |", "|---|---|---|---|"]
    for t in tools:
        drafts = [res[t][c["id"]] for c in cases if c["expect"] == "draft" and res[t].get(c["id"])]
        traps = [res[t][c["id"]] for c in cases if c["expect"] == "flag" and res[t].get(c["id"])]
        mean = round(sum(d["score"] for d in drafts) / len(drafts), 1) if drafts else "-"
        caught = (f"{sum(1 for d in traps if d['trap_caught'])}/{len(traps)}" if graded else "needs --grade") if traps else "-"
        lines.append(f"| {t} | {len(drafts)} | {mean} | {caught} |")

    lines += ["", "## Per document", "", "| Case | " + " | ".join(tools) + " |", "|---|" + "---|" * len(tools)]
    for c in cases:
        cells = []
        for t in tools:
            r = res[t].get(c["id"])
            if not r:
                cells.append("missing")
            elif c["expect"] == "flag":
                cells.append("caught" if r["trap_caught"] else ("missed" if graded else "?"))
            else:
                cells.append(str(r["score"]))
        lines.append(f"| {c['id']} | " + " | ".join(cells) + " |")

    lines += ["", "## Findings", ""]
    for c in cases:
        for t in tools:
            r = res[t].get(c["id"])
            if not r:
                continue
            det, avr = r["det"], r.get("averments")
            notes = []
            if avr and avr["missing"]:
                notes.append("Missing averments: " + "; ".join(avr["missing"]))
            if det["facts"]["missing"]:
                notes.append("Facts not reproduced: " + ", ".join(det["facts"]["missing"]))
            if det["old_law"]:
                notes.append("Old-law references: " + ", ".join(det["old_law"]))
            if det["placeholders"]:
                notes.append(f"Unfilled blanks: {det['placeholders']}")
            if det["unsupported_amounts"]:
                notes.append("Amounts not in the brief: " + ", ".join(det["unsupported_amounts"]))
            if det["diction_high"]:
                notes.append("Diction: " + "; ".join(det["diction_high"]))
            if det["citations"]:
                notes.append("Citations to verify: " + "; ".join(det["citations"]))
            if det["unlisted_dates"]:
                notes.append("Dates not in the brief (check derivation): " + ", ".join(det["unlisted_dates"]))
            if notes:
                lines.append(f"**{c['id']} · {t}**")
                lines += [f"- {n}" for n in notes]
                lines.append("")
    return "\n".join(lines)


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--outputs", required=True, type=Path, help="folder holding one sub-folder per tool")
    ap.add_argument("--grade", action="store_true", help="LLM-judge averments and traps (ANTHROPIC_API_KEY)")
    ap.add_argument("--only", help="comma-separated case ids")
    args = ap.parse_args(argv)

    cases = json.loads(CASES_FILE.read_text(encoding="utf-8"))["cases"]
    if args.only:
        wanted = set(args.only.split(","))
        cases = [c for c in cases if c["id"] in wanted]
    tools = sorted(p.name for p in args.outputs.iterdir() if p.is_dir())
    if not tools:
        print(f"No tool folders in {args.outputs}", file=sys.stderr)
        return 1

    res: Dict[str, Dict[str, Dict]] = {t: {} for t in tools}
    for t in tools:
        for c in cases:
            path = find_output(args.outputs / t, c["id"])
            if not path:
                continue
            text = load_text(path)
            det = deterministic(text, c)
            avr = judge(text, c) if args.grade else None
            res[t][c["id"]] = {
                "file": path.name,
                "det": det,
                "averments": avr,
                "score": overall(det, avr),
                "trap_caught": bool(avr and avr.get("defect_flagged")),
            }
            print(f"{t:>14} · {c['id']:<22} {res[t][c['id']]['score']}")

    (args.outputs / "results.json").write_text(json.dumps(res, indent=2, ensure_ascii=False), encoding="utf-8")
    report = render_report(cases, tools, res, args.grade)
    (args.outputs / "report.md").write_text(report, encoding="utf-8")
    print(f"\nWrote {args.outputs / 'report.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
