"""Refusal report — GENERATED FROM THE CODE, never hand-maintained.

Every drafting pipeline refuses (produces no document) rather than emit a defective
pleading when a legal precondition is unmet. This tool enumerates every refusal
condition and captures the EXACT message the advocate sees — by actually driving
each pipeline into the refusal and reading back its `reason`. Nothing here is a
transcribed copy of a message: if the code's wording changes, this report changes
with it, because it runs the code.

Each case is (trigger label, mutation) applied to a known-good baseline drawn from
that pipeline's own test suite (`tests/pipelines/test_<x>.py::_base`), so the
baseline can never drift from what the tests consider valid.

Usage:
    python tools/refusal_report.py            # print Markdown to stdout
    python tools/refusal_report.py -o out.md  # write Markdown
    python tools/refusal_report.py --docx out.docx   # write a .docx
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

# Make the repo root importable when run as a script.
_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.pipelines import (  # noqa: E402
    anticipatory_bail, consumer_complaint, default_bail, divorce_contested,
    divorce_mutual, dv_application, general_notice, injunction, maintenance,
    mact_claim, plaint, regular_bail, s138_complaint, s138_notice, summary_suit,
    writ_petition,
)
from src.rules import (  # noqa: E402
    bail_bar_detector as bars, compensation_multiplier as cm, custody_computation as cc,
    jurisdiction as jur,
)

# Baseline factories: reuse each pipeline's own test _base so a valid baseline is
# exactly what the tests treat as valid (guards against silent drift).
from tests.pipelines.test_general_notice import _base as base_general_notice  # noqa: E402
from tests.pipelines.test_s138_notice import _base as base_s138_notice  # noqa: E402
from tests.pipelines.test_s138_complaint import _base as base_s138_complaint  # noqa: E402
from tests.pipelines.test_regular_bail import _base as base_regular_bail  # noqa: E402
from tests.pipelines.test_anticipatory_bail import _base as base_anticipatory_bail  # noqa: E402
from tests.pipelines.test_default_bail import _base as base_default_bail  # noqa: E402
from tests.pipelines.test_injunction import _base as base_injunction  # noqa: E402
from tests.pipelines.test_plaint import _base as base_plaint  # noqa: E402
from tests.pipelines.test_summary_suit import _base as base_summary_suit  # noqa: E402
from tests.pipelines.test_writ_petition import _base as base_writ  # noqa: E402
from tests.pipelines.test_writ_petition import _private as writ_private_respondent  # noqa: E402
from tests.pipelines.test_maintenance import _base as base_maintenance  # noqa: E402
from tests.pipelines.test_dv_application import _base as base_dv  # noqa: E402
from tests.pipelines.test_consumer_complaint import _base as base_consumer  # noqa: E402
from tests.pipelines.test_mact_claim import _base as base_mact  # noqa: E402
from tests.pipelines.test_divorce_contested import _base as base_divorce_contested  # noqa: E402
from tests.pipelines.test_divorce_mutual import _first as base_mutual_first  # noqa: E402
from tests.pipelines.test_divorce_mutual import _second as base_mutual_second  # noqa: E402


@dataclass(frozen=True)
class Case:
    trigger: str
    mutation: Dict[str, Any]


@dataclass(frozen=True)
class DocSpec:
    spec: str
    title: str
    build: Callable[..., Any]
    base: Callable[..., Dict[str, Any]]
    cases: List[Case]


def _reason(result: Any) -> str:
    """Extract the refusal message the advocate sees, whatever the result shape."""
    reason = getattr(result, "reason", None)
    return reason if reason else "(no message captured)"


# --------------------------------------------------------------------------
# The refusal catalogue. Each mutation drives the pipeline into one refusal.
# --------------------------------------------------------------------------

def _catalogue() -> List[DocSpec]:
    return [
        DocSpec("2.1", "General demand notice", general_notice.build,
                base_general_notice, [
            Case("Non-positive sum demanded", {"demand_amount": 0}),
            Case("Compliance period below the statutory minimum",
                 {"compliance_days": 7, "statutory_minimum_days": 15}),
            Case("No specific consequence of non-compliance stated",
                 {"civil_consequence": "", "criminal_consequence": None}),
            Case("Dispatch-proof placeholder blank",
                 {"dispatch_tracking_placeholder": "  "}),
        ]),
        DocSpec("2.2", "s.138 statutory notice", s138_notice.build,
                base_s138_notice, [
            Case("No cheque particulars", {"cheques": []}),
            Case("Memo receipt date missing", {"memo_received_date": None}),
            Case("Dishonour reason blank", {"dishonour_reason": "  "}),
            Case("Dispatch-proof placeholder blank",
                 {"dispatch_tracking_placeholder": "  "}),
            Case("Compliance period not exactly 15 days", {"compliance_days": 20}),
        ]),
        DocSpec("3.1", "s.138 complaint", s138_complaint.build,
                base_s138_complaint, [
            Case("Complainant not the payee / holder in due course",
                 {"is_payee": False}),
            Case("Company accused without the s.141 averment",
                 {"accused_is_company": True, "s141_averment": None}),
            Case("Payee's bank branch outside the court's jurisdiction",
                 {"payee_bank_branch_in_jurisdiction": False}),
        ]),
        DocSpec("3.2", "Regular bail", regular_bail.build, base_regular_bail, [
            Case("Special-statute bail bar not addressed",
                 {"special_statute": bars.Statute.NDPS, "bail_bar_addressed": False}),
            Case("Successive application without a prior-order reference",
                 {"is_successive": True, "prior_order_ref": None}),
            Case("Successive application without a change-of-circumstance averment",
                 {"is_successive": True, "prior_order_ref": "Crl.OP 1/2024",
                  "change_of_circumstance": None}),
        ]),
        DocSpec("3.3", "Anticipatory bail", anticipatory_bail.build,
                base_anticipatory_bail, [
            Case("Applicant already arrested (routed to regular bail)",
                 {"already_arrested": True}),
            Case("SC/ST offence — s.18/18A bar",
                 {"already_arrested": False, "is_scst_offence": True}),
            Case("Reasonable apprehension of arrest not stated",
                 {"apprehension_basis": "  "}),
        ]),
        DocSpec("3.4", "Default bail", default_bail.build, base_default_bail, [
            Case("Readiness to furnish bail not averred",
                 {"readiness_to_furnish_bail": False}),
            Case("Chargesheet preceded the application (right extinguished)",
                 {"application_ts": datetime(2024, 6, 2, 12, 0),
                  "chargesheet_filed_ts": datetime(2024, 6, 1, 12, 0)}),
            Case("Precedence indeterminate (equal timestamps)",
                 {"application_ts": datetime(2024, 6, 1, 12, 0),
                  "chargesheet_filed_ts": datetime(2024, 6, 1, 12, 0)}),
        ]),
        DocSpec("4.1", "Plaint (money recovery)", plaint.build, base_plaint, [
            Case("Cause of action missing a date", {"cause_of_action_date": None}),
            Case("Cause of action missing a place", {"cause_of_action_place": "  "}),
            Case("No territorial jurisdiction basis pleaded",
                 {"territorial_bases": []}),
            Case("Government defendant without an s.80 CPC notice averment",
                 {"is_government_defendant": True, "s80_notice_ref": None}),
            Case("Commercial suit without a s.12A certificate/exemption",
                 {"is_commercial_suit": True, "s12a_ref": None}),
        ]),
        DocSpec("4.1-O37", "Summary suit (O37)", summary_suit.build,
                base_summary_suit, [
            Case("Claim outside the scope of Order XXXVII",
                 {"claim_basis": summary_suit.ClaimBasis.OTHER}),
        ]),
        DocSpec("4.2", "Temporary injunction", injunction.build, base_injunction, [
            Case("Prima facie case limb missing", {"prima_facie_case": ""}),
            Case("Balance of convenience limb missing",
                 {"balance_of_convenience": ""}),
            Case("Irreparable injury limb missing", {"irreparable_injury": ""}),
        ]),
        DocSpec("5.1", "Writ petition (Art. 226)", writ_petition.build, base_writ, [
            Case("Mandamus without a legal right", {"legal_right": ""}),
            Case("No State / public-function respondent",
                 {"_respondents_private": True}),
            Case("Alternative remedy not addressed",
                 {"alternative_remedy_addressed": False}),
        ]),
        DocSpec("6.1", "Contested divorce (HMA s.13)", divorce_contested.build,
                base_divorce_contested, [
            Case("Cruelty pleaded without a dated instance",
                 {"ground": divorce_contested.Ground.CRUELTY,
                  "cruelty_instances": []}),
        ]),
        DocSpec("6.2-first", "Mutual-consent divorce — first motion",
                divorce_mutual.build_first_motion, base_mutual_first, [
            Case("Separation under one year",
                 {"living_separately_since": date(2024, 1, 1),
                  "first_motion_date": date(2024, 6, 1)}),
        ]),
        DocSpec("6.2-second", "Mutual-consent divorce — second motion",
                divorce_mutual.build_second_motion, base_mutual_second, [
            Case("Second motion after eighteen months",
                 {"first_motion_order_date": date(2022, 1, 1),
                  "second_motion_date": date(2024, 6, 1)}),
        ]),
        DocSpec("6.3", "Maintenance (BNSS 144 / CrPC 125)", maintenance.build,
                base_maintenance, [
            Case("Neglect or refusal to maintain not pleaded",
                 {"neglect_or_refusal_pleaded": False}),
            Case("Petitioner's Rajnesh affidavit absent",
                 {"petitioner_rajnesh_affidavit": False}),
        ]),
        DocSpec("6.4", "DV Act s.12 application", dv_application.build, base_dv, [
            Case("Domestic relationship not pleaded",
                 {"domestic_relationship_pleaded": False}),
        ]),
        DocSpec("7.1", "Consumer complaint", consumer_complaint.build,
                base_consumer, [
            Case("Goods/services obtained for a commercial purpose",
                 {"purpose": consumer_complaint.Purpose.COMMERCIAL}),
            Case("Affidavit absent", {"affidavit_present": False}),
            Case("Evidence affidavit absent",
                 {"evidence_affidavit_present": False}),
        ]),
        DocSpec("7.2", "MACT claim (MV Act s.166)", mact_claim.build, base_mact, [
            Case("Rash and negligent driving not pleaded",
                 {"rash_negligent_pleaded": False}),
        ]),
    ]


def _run_case(spec: DocSpec, case: Case):
    """Run one refusal case; return the pipeline result (or a raised exception)."""
    mutation = dict(case.mutation)
    # A couple of cases need to rebuild structured inputs rather than swap a scalar.
    if mutation.pop("_respondents_private", False):
        kwargs = spec.base()
        kwargs["respondents"] = [writ_private_respondent(serial=1)]
    else:
        kwargs = spec.base(**mutation)
    return spec.build(**kwargs)


def _apply(spec: DocSpec, case: Case) -> str:
    """Run one refusal case and return the captured message."""
    try:
        result = _run_case(spec, case)
    except Exception as exc:  # a raised input error is also a refusal-to-generate
        return f"{type(exc).__name__}: {exc}"
    if getattr(result, "refused", False):
        return _reason(result)
    return "(did NOT refuse — case needs review)"


def _capture(spec: DocSpec, case: Case):
    """Run a case and return its structured Refusal (or None if it raised / did not
    refuse / carries no coded Refusal)."""
    try:
        result = _run_case(spec, case)
    except Exception:
        return None
    if not getattr(result, "refused", False):
        return None
    return getattr(result, "refusal", None)


def build_report() -> str:
    """Render the refusal report as Markdown, generated by running the pipelines.

    This is Section C AND the typed code catalogue: every refusal's stable `code`,
    `severity`, the implicated fields, and the exact advocate-facing message — all
    captured by driving the real pipeline, so it cannot drift from the engine.
    """
    specs = _catalogue()
    rows = 0
    uncoded = 0
    out: List[str] = []
    out.append("# Section C — Refusal Conditions (typed catalogue)")
    out.append("")
    out.append(
        "Every condition under which a pipeline refuses to generate a document, "
        "with its stable machine code, severity, the implicated fields, and the "
        "exact message the advocate sees. GENERATED by driving each pipeline into "
        "the refusal and reading back its structured Refusal — not hand-maintained, "
        "so it cannot drift from the code. The `code` column is the contract the "
        "frontend branches on."
    )
    out.append("")

    def esc(s: str) -> str:
        return str(s).replace("|", "\\|").replace("\n", " ")

    for spec in specs:
        out.append(f"## {spec.spec} — {spec.title}")
        out.append("")
        out.append("| Code | Sev | Fields | Message the advocate sees |")
        out.append("|---|---|---|---|")
        for case in spec.cases:
            rf = _capture(spec, case)
            message = esc(_apply(spec, case))
            if rf is None:
                uncoded += 1
                out.append(f"| _(uncoded)_ | — | — | {message} |")
            else:
                fields = ", ".join(rf.fields) or "—"
                out.append(f"| `{rf.code}` | {rf.severity.value} | {esc(fields)} "
                           f"| {message} |")
            rows += 1
        out.append("")

    summary = f"\n_{rows} refusal conditions across {len(specs)} documents"
    summary += f"; all coded._" if uncoded == 0 else f"; {uncoded} UNCODED._"
    out.insert(3, summary)
    return "\n".join(out)


def _write_docx(markdown_text: str, path: Path) -> None:
    """Render the report to a simple .docx (headings, paragraphs, table rows)."""
    from docx import Document
    doc = Document()
    for line in markdown_text.splitlines():
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("|") and set(line) <= set("|-: "):
            continue  # table separator row
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            doc.add_paragraph(" — ".join(cells))
        elif line.strip():
            doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Generate the refusal report from code.")
    ap.add_argument("-o", "--out", type=Path, help="write Markdown to this path")
    ap.add_argument("--docx", type=Path, help="write a .docx to this path")
    args = ap.parse_args(argv)

    report = build_report()
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(report, encoding="utf-8")
        print(f"wrote {args.out}")
    if args.docx:
        _write_docx(report, args.docx)
        print(f"wrote {args.docx}")
    if not args.out and not args.docx:
        sys.stdout.reconfigure(encoding="utf-8")
        print(report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
