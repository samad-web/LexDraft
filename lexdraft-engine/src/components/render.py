"""docxtpl rendering + text extraction helpers for shared components.

Templates are real .docx files under templates/shared/, authored with Jinja-style
placeholders. Output is always editable .docx (never PDF, never plain text).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping, Union

from docxtpl import DocxTemplate
from docx import Document

# Repo root -> templates/shared
_TEMPLATES_DIR = Path(__file__).resolve().parents[2] / "templates" / "shared"
# Imported formats live apart from the shared set (src/imports/manifest.py).
_IMPORTED_DIR = Path(__file__).resolve().parents[2] / "templates" / "imported"
_IMPORTED_PREFIX = "imported/"


def template_path(name: str) -> Path:
    """Resolve a shared-template filename to its absolute path. An
    'imported/<file>.docx' name resolves under templates/imported/."""
    if name.startswith(_IMPORTED_PREFIX):
        return _IMPORTED_DIR / name[len(_IMPORTED_PREFIX):]
    return _TEMPLATES_DIR / name


def render_to_docx(
    template: Union[str, Path],
    context: Mapping[str, Any],
    out_path: Union[str, Path],
) -> Path:
    """Render `template` with `context` to a .docx at `out_path`."""
    tpl_path = template if isinstance(template, Path) else template_path(str(template))
    if not Path(tpl_path).exists():
        raise FileNotFoundError(
            f"template not found: {tpl_path}. Build templates first "
            f"(python templates/build_templates.py)."
        )
    doc = DocxTemplate(str(tpl_path))
    doc.render(dict(context))
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def extract_text(docx_path: Union[str, Path]) -> str:
    """Extract paragraph text from a .docx, one paragraph per line.

    Deterministic: only the text and its order, which is the legally meaningful
    content golden tests pin. Trailing whitespace per line is stripped; blank
    paragraphs are preserved as empty lines so structure is visible.
    """
    doc = Document(str(docx_path))
    lines = [p.text.rstrip() for p in doc.paragraphs]
    return "\n".join(lines).strip() + "\n"
