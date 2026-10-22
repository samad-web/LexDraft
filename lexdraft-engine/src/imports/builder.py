"""Format spec -> .docx template (+ validation and a smoke render).

A *format spec* is the structured transcription of a source format (produced by
src/imports/transcribe.py, saved under templates/imported/specs/<slug>.json so a
template can be rebuilt or hand-corrected). Its `blocks` are paragraphs in
document order; placeholders are `{{ field_key }}` and are written into ONE run
per paragraph, so docxtpl never meets a split placeholder.

Page setup follows docs/template-authoring-guide.md: A4, 1.5" left binding
margin, 1" elsewhere, Carlito body / Caladea headings, "Page X of Y" header.
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import manifest as mf

BLOCK_STYLES = ("title", "heading", "center", "para", "left", "right", "numbered",
                "signature", "note", "blank", "list")

_PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")
_CONTROL = re.compile(r"\{%|%\}|\{#|#\}")
_STRAY = re.compile(r"\{\{|\}\}")

BODY_FONT = "Carlito"
HEADING_FONT = "Caladea"


# --------------------------------------------------------------------------
# Normalisation + validation
# --------------------------------------------------------------------------

def normalise_spec(spec: Dict[str, Any]) -> Dict[str, Any]:
    """Bring a transcribed spec to the stored shape (idempotent): enum_values only
    on enum fields, string_list samples as lists (the transcriber sends "a | b")."""
    for f in spec.get("fields") or []:
        if not isinstance(f, dict):
            continue
        if f.get("kind") != "enum":
            f.pop("enum_values", None)
        if f.get("kind") == "string_list" and isinstance(f.get("sample"), str):
            f["sample"] = [s.strip() for s in f["sample"].split("|") if s.strip()]
    return spec


def check_spec(spec: Mapping[str, Any]) -> List[str]:
    """Every problem that would make the spec render wrongly (empty = OK)."""
    problems: List[str] = []
    fields = spec.get("fields") or []
    kinds = {f.get("key"): f.get("kind") for f in fields if isinstance(f, Mapping)}
    entry_like = {"title": spec.get("title"), "status": "pending_review",
                  "template": "x.docx", "fields": fields}
    problems += mf.validate_entry(spec.get("slug", ""), entry_like)

    blocks = spec.get("blocks") or []
    if not blocks:
        problems.append("blocks must be a non-empty list")
    used = set()
    for i, b in enumerate(blocks):
        where = f"blocks[{i}]"
        style, text = b.get("style"), b.get("text") or ""
        if style not in BLOCK_STYLES:
            problems.append(f"{where}.style {style!r} is not one of {BLOCK_STYLES}")
            continue
        if style == "list":
            if kinds.get(text.strip()) != "string_list":
                problems.append(f"{where}: a list block's text must name a string_list field")
            used.add(text.strip())
            continue
        if _CONTROL.search(text):
            problems.append(f"{where}: template control tags are not allowed in text")
        for key in _PLACEHOLDER.findall(text):
            used.add(key)
            if key not in kinds:
                problems.append(f"{where}: placeholder {{{{ {key} }}}} has no declared field")
            elif kinds[key] == "string_list":
                problems.append(f"{where}: string_list field {key!r} must use a list block")
        if _STRAY.search(_PLACEHOLDER.sub("", text)):
            problems.append(f"{where}: malformed placeholder braces")
        cond = (b.get("condition") or "").strip()
        if cond:
            if kinds.get(cond) != "bool":
                problems.append(f"{where}.condition {cond!r} must name a bool field")
            used.add(cond)
    for key in kinds:
        if key not in used:
            problems.append(f"field {key!r} is declared but never placed in the template")
    return problems


# --------------------------------------------------------------------------
# Rendering the template
# --------------------------------------------------------------------------

def _set_font(style, name: str, size: float, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _field_run(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(code)
    run._r.append(end)


def _page_setup(doc) -> None:
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = Cm(3.81)
        section.right_margin = section.top_margin = section.bottom_margin = Cm(2.54)
        section.header_distance = section.footer_distance = Cm(1.27)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.add_run("Page ")
        _field_run(header, "PAGE")
        header.add_run(" of ")
        _field_run(header, "NUMPAGES")
    _set_font(doc.styles["Normal"], BODY_FONT, 12)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)


def _normalise(text: str) -> str:
    return _PLACEHOLDER.sub(lambda m: "{{ " + m.group(1) + " }}", text)


def _add_block(doc, block: Mapping[str, Any]) -> None:
    style = block["style"]
    text = _normalise(block.get("text") or "")
    if style == "blank":
        doc.add_paragraph()
        return
    if style == "list":
        key = text.strip()
        doc.add_paragraph("{%p for item in " + key + " %}")
        doc.add_paragraph("{{ item }}", style="List Number")
        doc.add_paragraph("{%p endfor %}")
        return
    cond = (block.get("condition") or "").strip()
    if cond:
        text = "{% if " + cond + " %}" + text + "{% endif %}"

    p = doc.add_paragraph(style="List Number" if style == "numbered" else None)
    run = p.add_run(text)  # one run: placeholders can't split
    run.bold = bool(block.get("bold")) or style in ("title", "heading")
    fmt = p.paragraph_format
    if style in ("title", "heading", "center"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style in ("title", "heading"):
            run.font.name = HEADING_FONT
            run.font.size = Pt(14 if style == "title" else 12)
    elif style in ("right", "signature"):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if style == "signature":
            fmt.space_before = Pt(18)
    elif style == "para":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.first_line_indent = Cm(1.27)
    elif style == "numbered":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif style == "note":
        run.font.size = Pt(10)


def build_docx(spec: Mapping[str, Any], out_path: Path) -> Path:
    """Write the template for `spec` to `out_path` (caller validated it)."""
    doc = Document()
    _page_setup(doc)
    # Rule 6 banner + imported/unverified marker + court-issued marker: the same
    # three lines every generated template carries (templates/build_templates.py).
    doc.add_paragraph("{% if review_banner %}{{ review_banner }}{% endif %}")
    doc.add_paragraph("{% if unverified_marker %}{{ unverified_marker }}{% endif %}")
    if spec.get("court_issued"):
        doc.add_paragraph("{% if court_issues %}[COURT ISSUES THIS — the advocate prepares "
                          "this draft for the court's issue; it is not an advocate-authored "
                          "pleading]{% endif %}")
    for block in spec["blocks"]:
        _add_block(doc, block)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------
# Smoke render — the template renders through the real adapter + renderer
# --------------------------------------------------------------------------

def smoke_render(slug: str, entry: Mapping[str, Any], template_file: Path) -> List[str]:
    """Render `template_file` with the entry's sample payload through the generic
    adapter and the production renderer; return problems (empty = clean)."""
    from ..components.render import extract_text, render_to_docx
    from ..service.adapters.imported import adapter_for, sample_payload

    adapter = adapter_for(slug, entry)
    try:
        kwargs = adapter.coerce(sample_payload(entry))
        result = adapter.build(**kwargs)
        ctx = dict(result.context)
        ctx["review_banner"] = "[DRAFT — FOR ADVOCATE REVIEW]"
        with tempfile.TemporaryDirectory() as tmp:
            out = render_to_docx(template_file, ctx, Path(tmp) / "smoke.docx")
            text = extract_text(out)
    except Exception as exc:  # any failure here is a broken template
        return [f"smoke render failed: {type(exc).__name__}: {exc}"]
    leftovers = re.findall(r"\{\{.*?\}\}|\{%.*?%\}", text)
    return [f"unrendered template text in output: {t}" for t in leftovers]
