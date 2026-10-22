"""Source format (PDF / scan / photo / Word / text) -> format spec, via Claude.

The model reads the source format and returns a structured spec (JSON schema
enforced by structured outputs): the fields to collect, the paragraphs with
`{{ field }}` placeholders where the blanks were, statutory-reference updates,
corrections, and anything it could not read with confidence. Nothing here is
trusted blindly — builder.check_spec validates the spec and the engine marks the
format "not yet approved" until an advocate approves it.

Needs the `import` extra (pip install -e ".[import]") and ANTHROPIC_API_KEY
(read from the environment, else from ../LexDraft/apps/api/.env).
"""

from __future__ import annotations

import base64
import json
import os
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from .builder import BLOCK_STYLES, normalise_spec
from .manifest import FIELD_KINDS, ROOT

DEFAULT_MODEL = "claude-opus-5-5"

_MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}

SPEC_SCHEMA: Dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": ["is_legal_format", "equivalent_existing_doc_type", "slug", "title",
                 "court_issued", "fields", "blocks", "law_updates", "corrections",
                 "uncertain", "summary"],
    "properties": {
        "is_legal_format": {"type": "boolean"},
        "equivalent_existing_doc_type": {"type": "string"},
        "slug": {"type": "string"},
        "title": {"type": "string"},
        "court_issued": {"type": "boolean"},
        "fields": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["key", "kind", "required", "label", "help", "enum_values", "sample"],
                "properties": {
                    "key": {"type": "string"},
                    "kind": {"type": "string", "enum": list(FIELD_KINDS)},
                    "required": {"type": "boolean"},
                    "label": {"type": "string"},
                    "help": {"type": "string"},
                    "enum_values": {"type": "array", "items": {"type": "string"}},
                    "sample": {"type": "string"},
                },
            },
        },
        "blocks": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["style", "text", "bold", "condition"],
                "properties": {
                    "style": {"type": "string", "enum": list(BLOCK_STYLES)},
                    "text": {"type": "string"},
                    "bold": {"type": "boolean"},
                    "condition": {"type": "string"},
                },
            },
        },
        "law_updates": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["original", "updated", "note"],
                "properties": {"original": {"type": "string"}, "updated": {"type": "string"},
                               "note": {"type": "string"}},
            },
        },
        "corrections": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["original", "corrected"],
                "properties": {"original": {"type": "string"}, "corrected": {"type": "string"}},
            },
        },
        "uncertain": {"type": "array", "items": {"type": "string"}},
        "summary": {"type": "string"},
    },
}

SYSTEM_PROMPT = """\
You convert an Indian court / legal drafting FORMAT into a structured template spec \
for LexDraft, a drafting engine used by practising advocates. The spec you return is \
turned into a Word template; the engine fills its placeholders from the matter's facts. \
An advocate reviews every imported format before it is relied on, so flag anything \
you are not sure about instead of guessing.

How to transcribe:
- Reproduce the format faithfully: same headings, cause-title layout, recitals, \
paragraph order and operative wording. Do not add legal content the format does not \
have, and do not drop any part of it.
- Every blank, dotted line, "____", "the ___ day of ___ 20__", or bracketed \
alternative the drafter fills in becomes a `{{ field_key }}` placeholder. A full \
date ("the __ day of __ 20__") is ONE field of kind "date". Case numbers such as \
"O.S. No. __ of 20__" become one string field (e.g. case_ref).
- Court names and places that vary between matters become fields (court_name, \
court_place); keep fixed statutory wording as text. Alternatives like \
"Plaintiff / Petitioner" may stay as printed text.
- Fix obvious typing/OCR errors of the source and list each one in `corrections`.
- Leave out publisher branding, watermarks (e.g. "LAWE"), page numbers and scan noise.
- Where the format cites a REPEALED enactment, write the current provision \
(CrPC 1973 -> Bharatiya Nagarik Suraksha Sanhita 2023; IPC -> Bharatiya Nyaya Sanhita \
2023; Indian Evidence Act 1872 -> Bharatiya Sakshya Adhiniyam 2023) ONLY when you are \
certain of the corresponding section, and record every change in `law_updates`. If you \
are not certain, keep the original citation and say so in `uncertain`. Civil \
Procedure Code references stay as they are.

Fields:
- `key`: snake_case. REUSE a key from the common-keys list whenever it means the same \
thing (the advocate's matter data already uses those names); otherwise invent a clear one.
- `kind`: string | date | int | bool | enum | string_list. Use enum (with \
`enum_values`) only for a genuine fixed choice; string_list for a repeatable list \
(e.g. schedule of documents). `enum_values` is [] for every other kind.
- `required`: true only for facts without which the document is meaningless (party \
names, case number, the core particulars). Fill-ins a drafter often leaves for hand \
completion are required: false; an empty optional field prints as a blank line.
- `label`: the human question label; `help`: one short hint (may be ""); `sample`: a \
realistic sample value as text ("" if none; ISO YYYY-MM-DD for dates; "true"/"false" \
for bool; for string_list, items separated by " | ").

Blocks (paragraphs, in document order):
- `style`: title (document heading), heading (section heading such as SCHEDULE / \
PRAYER), center, para (justified body paragraph), left, right, numbered (one item of a \
numbered list — do NOT type the number), signature (a right-aligned signature / \
designation line), note (small print, e.g. N.B. notes), blank (an empty line; text ""), \
list (renders a string_list field as numbered items; `text` is just that field's key).
- `text` may contain only literal text and `{{ field_key }}` placeholders — no other \
template syntax. Every declared field must appear in at least one block.
- `condition`: "" normally; or the key of a bool field when the paragraph should print \
only if that bool is true.
- Tables: linearise them into paragraphs (one row per paragraph, "Label: {{ key }}").

Classification:
- `is_legal_format`: false if the input is not a legal drafting format at all.
- `equivalent_existing_doc_type`: if this format is substantively the SAME document \
(same purpose and structure, differing only in court name, place or trivial wording) as \
one in the existing list, give that slug; otherwise "". A different document type of the \
same family (e.g. a different kind of summons) is NOT equivalent.
- `court_issued`: true for instruments the court issues under its seal (summons, court \
notices, warrants, orders), false for documents a party/advocate files or executes.
- `slug`: snake_case doc-type id (3-61 chars) not in the existing list; `title`: a short \
human title.
- `summary`: two sentences on what the format is and when it is used."""


def resolve_api_key() -> Optional[str]:
    """ANTHROPIC_API_KEY from the environment, else from the monorepo API .env."""
    key = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if key:
        return key
    for env_path in (ROOT.parent / "LexDraft" / "apps" / "api" / ".env",
                     ROOT.parent / "apps" / "api" / ".env"):
        if not env_path.exists():
            continue
        for raw in env_path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if line.startswith("ANTHROPIC_API_KEY="):
                return line.split("=", 1)[1].strip().strip("'\"") or None
    return None


def _docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def source_content(path: Path) -> List[Dict[str, Any]]:
    """The user-message content blocks that carry the source format."""
    suffix = path.suffix.lower()
    if suffix in _MEDIA_TYPES:
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        kind = "document" if suffix == ".pdf" else "image"
        return [{"type": kind,
                 "source": {"type": "base64", "media_type": _MEDIA_TYPES[suffix], "data": data}}]
    if suffix == ".docx":
        text = _docx_text(path)
    elif suffix in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"unsupported source format {suffix!r} "
                         f"(use PDF, PNG/JPG/WEBP, DOCX or TXT)")
    return [{"type": "text", "text": "SOURCE FORMAT (text of the document):\n\n" + text}]


def _instructions(existing: Iterable[Tuple[str, str]], common_keys: Iterable[str],
                  hint_title: str) -> str:
    listing = "\n".join(f"- {slug}: {title}" for slug, title in sorted(existing))
    hint = f"\nThe source is titled: {hint_title}\n" if hint_title else ""
    return (f"{hint}\nExisting doc types (slug: title):\n{listing}\n\n"
            f"Common field keys to reuse: {', '.join(sorted(common_keys))}\n\n"
            "Transcribe the source format above into the spec.")


def transcribe(path: Path, *, existing: Iterable[Tuple[str, str]],
               common_keys: Iterable[str], hint_title: str = "",
               model: str = DEFAULT_MODEL) -> Dict[str, Any]:
    """Call Claude on the source format at `path`; return the parsed spec."""
    import anthropic

    api_key = resolve_api_key()
    if not api_key:
        raise RuntimeError("ANTHROPIC_API_KEY is not set (env or LexDraft/apps/api/.env)")
    client = anthropic.Anthropic(api_key=api_key)
    content = source_content(path) + [
        {"type": "text", "text": _instructions(existing, common_keys, hint_title)}]
    with client.beta.messages.stream(
        model=model,
        max_tokens=64000,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        system=SYSTEM_PROMPT,
        output_config={"effort": "high",
                       "format": {"type": "json_schema", "schema": SPEC_SCHEMA}},
        messages=[{"role": "user", "content": content}],
    ) as stream:
        message = stream.get_final_message()

    if message.stop_reason == "refusal":
        raise RuntimeError(f"the model declined to transcribe {path.name}")
    if message.stop_reason == "max_tokens":
        raise RuntimeError(f"transcription of {path.name} was cut off (max_tokens)")
    text = next(b.text for b in message.content if b.type == "text")
    return normalise_spec(json.loads(text))
