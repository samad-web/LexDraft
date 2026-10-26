"""Write an advocate's text edits back INTO an engine-rendered .docx.

The Draft page shows a draft as one text block per non-empty paragraph (see
app._extract_blocks). When the advocate edits that text and downloads, a generic
HTML export would lose the court layout (alignment, headings, tab leaders, page
numbers). Instead, map the edited paragraphs back onto the document's own
paragraphs, in order, and replace only the ones whose text changed — keeping
each paragraph's style, alignment and its first run's formatting.

Only possible when the paragraph count still matches (the advocate edited text
in place, e.g. filled in blanks). If paragraphs were added or removed, return
None and let the caller fall back to its generic export.
"""

from __future__ import annotations

import copy
import io
import re
from typing import List, Optional

from docx import Document
from docx.text.paragraph import Paragraph


def _targets(doc) -> List[tuple]:
    """(paragraph, line_count) for every paragraph with text, in reading order —
    each line of a paragraph (soft line breaks) is one display block."""
    out = []
    for p in doc.paragraphs:
        lines = [l for l in p.text.split("\n") if l.strip()]
        if lines:
            out.append((p, len(lines)))
    return out


def apply_edits(docx_bytes: bytes, paragraphs: List[str]) -> Optional[bytes]:
    doc = Document(io.BytesIO(docx_bytes))
    targets = _targets(doc)
    edited = [t.rstrip() for t in paragraphs if t.strip()]
    if sum(n for _, n in targets) != len(edited):
        return None
    i = 0
    for p, n in targets:
        new_text = "\n".join(edited[i:i + n])
        i += n
        old_text = "\n".join(l for l in p.text.split("\n") if l.strip())
        if new_text == old_text.rstrip():
            continue  # untouched paragraph keeps every run exactly as rendered
        runs = p.runs
        if not runs:
            p.add_run(new_text)
            continue
        runs[0].text = new_text  # python-docx turns \n / \t into breaks / tabs
        for r in runs[1:]:
            r.text = ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_LEADING_NUMBER = re.compile(r"^\s*(?:\d+|[ivxlc]+|[a-z])[.)]\s+", re.IGNORECASE)


def _set_text(p, text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def _auto_numbered(p) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.numPr is not None


def splice_lines(docx_bytes: bytes, start: int, end: int, paragraphs: List[str]) -> Optional[bytes]:
    """Replace display lines [start, end] (inclusive, the same line indexing as the
    Draft page's blocks) with `paragraphs` — one Word paragraph each — keeping the
    court layout everywhere else. Used to write a COMPOSED body back into the engine
    .docx: the cause title above and the prayer/verification below are untouched.

    Reused paragraphs keep their own style and first-run formatting; extra ones are
    clones of the last paragraph in the range; surplus ones are removed. A paragraph
    Word numbers automatically has any typed "1." prefix stripped so it isn't
    numbered twice. Returns None when the range is out of bounds or cuts through a
    multi-line paragraph — the caller then keeps the plain engine draft.
    """
    doc = Document(io.BytesIO(docx_bytes))
    targets = _targets(doc)
    new = [t.strip() for t in paragraphs if t.strip()]
    total = sum(n for _, n in targets)
    if not new or start < 0 or end < start or end >= total:
        return None

    # Map the line range onto whole paragraphs.
    picked, line = [], 0
    for p, n in targets:
        first, last = line, line + n - 1
        line += n
        if last < start or first > end:
            continue
        if first < start or last > end:
            return None  # range splits a soft-broken paragraph
        picked.append(p)
    if not picked:
        return None

    model, insert_after = picked[-1], picked[-1]._p
    for i, text in enumerate(new):
        if i < len(picked):
            p = picked[i]
        else:
            clone = copy.deepcopy(model._p)
            insert_after.addnext(clone)
            insert_after = clone
            p = Paragraph(clone, model._parent)
        _set_text(p, _LEADING_NUMBER.sub("", text) if _auto_numbered(p) else text)
    for p in picked[len(new):]:
        p._p.getparent().remove(p._p)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
