"""Blank forms — any doc type rendered with fill-in lines instead of facts.

For the Draft page's "Download blank (Word)": the advocate gets the format itself,
editable, to fill by hand or in Word. Built through the SAME adapter + template as
a real draft, so the layout is identical:

  * every REQUIRED field gets a stand-in (a fill-in line for text; sentinel values
    for dates / numbers, swapped back to fill-in lines after rendering); optional
    fields stay absent, so optional paragraphs are left out as in a bare form;
  * an enum takes the caller's `preset` (e.g. a warrant's kind) or its first value;
  * if the pipeline still refuses (a format whose legal checks need real facts,
    e.g. a s.138 notice's limitation window), the caller gets that reason and the
    UI offers "Draft this" instead.

No review banner: a blank form carries no generated content to review.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Mapping, Optional

from docx import Document

from .adapters import Adapter, AdapterValidationError, FieldSpec, ListFieldSpec

BLANK = "____________"
_DATE_SENTINEL = "1900-01-01"
_INT_SENTINEL = 99999

# The date sentinel as pipelines may print it (ISO, d/m/y, long form).
_SENTINEL_RE = re.compile(
    r"\b1900-01-01\b"
    r"|\b0?1(?:st)?[ ./-]0?1[ ./-]1900\b"
    r"|\b(?:the\s+)?0?1(?:st)?\s+(?:day\s+of\s+)?January,?\s+1900\b"
    r"|\bJanuary\s+0?1(?:st)?,?\s+1900\b"
    r"|\b99,?999\b"
    r"|\b1900\b",  # a year derived from the date stand-in ("Suit No. __ of 1900")
    re.IGNORECASE,
)


class BlankRefused(Exception):
    """The format can't be produced without real facts (a legal hard-stop)."""


def _stand_in(spec: FieldSpec, preset: Mapping[str, Any]) -> Any:
    if spec.key in preset:
        return preset[spec.key]
    kind = spec.kind
    if kind == "enum":
        return (spec.enum_values or ("",))[0]
    if kind == "date":
        return _DATE_SENTINEL
    if kind == "int":
        return _INT_SENTINEL
    if kind == "bool":
        return True   # averments print, as on the printed form
    if kind == "string_list":
        return [BLANK]
    return BLANK


def blank_payload(adapter: Adapter, preset: Optional[Mapping[str, Any]] = None) -> Dict[str, Any]:
    """An inbound payload that satisfies every required field with a stand-in."""
    preset = dict(preset or {})
    out: Dict[str, Any] = {}
    for spec in adapter.fields:
        if isinstance(spec, ListFieldSpec):
            if spec.required or spec.key in preset:
                item = {sub.key: _stand_in(sub, {}) for sub in spec.item_fields if sub.required}
                out[spec.key] = preset.get(spec.key) or [item] * max(1, spec.min_items)
            continue
        if spec.required or spec.key in preset:
            out[spec.key] = _stand_in(spec, preset)
    return out


def blank_context(adapter: Adapter, preset: Optional[Mapping[str, Any]] = None) -> Dict[str, Any]:
    """Build the doc with stand-ins; raise BlankRefused if the pipeline refuses."""
    try:
        kwargs = adapter.coerce(blank_payload(adapter, preset))
        result = adapter.build(**kwargs)
    except (AdapterValidationError, ValueError, TypeError, KeyError, AttributeError) as exc:
        # A field whose value feeds a typed engine object (an enum list, a computed
        # date window) can't take a stand-in: this format needs real facts.
        raise BlankRefused("this format needs the case details to be drafted") from exc
    if getattr(result, "refused", False):
        refusal = getattr(result, "refusal", None)
        raise BlankRefused(getattr(refusal, "reason", "") or getattr(result, "reason", "")
                           or "this format needs the case details to be drafted")
    ctx = dict(result.context)
    ctx["review_banner"] = ""
    ctx["unverified_marker"] = ""
    return ctx


def scrub_sentinels(docx_path: Path) -> None:
    """Swap any rendered sentinel date / number back to a fill-in line, in place."""
    doc = Document(str(docx_path))

    def paragraphs():
        yield from doc.paragraphs
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for section in doc.sections:
            yield from section.header.paragraphs
            yield from section.footer.paragraphs

    changed = False
    for p in paragraphs():
        if not _SENTINEL_RE.search(p.text):
            continue
        hit = False
        for run in p.runs:
            if _SENTINEL_RE.search(run.text):
                run.text = _SENTINEL_RE.sub(BLANK, run.text)
                hit = True
        if not hit and p.runs:  # sentinel split across runs: collapse into the first
            text = _SENTINEL_RE.sub(BLANK, p.text)
            p.runs[0].text = text
            for run in p.runs[1:]:
                run.text = ""
        changed = True
    if changed:
        doc.save(str(docx_path))
