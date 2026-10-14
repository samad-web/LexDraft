"""Build the GOLD-STANDARD s138_notice.docx — the hand-authored reference for the
s.138 statutory demand notice.

A s.138 notice is a LETTER (advocate correspondence), not a court pleading: no cause
title, no verification/prayer. It carries the advocate's reference line, the addressee
block, a subject line, numbered factual recitals, a cheque-particulars list (one entry
per cheque — the whole point of the multi-cheque support), the statutory 15-day demand,
the consequence of non-compliance, and the advocate signature.

Real OOXML court-correspondence formatting (what python-docx defaults lack):
  * A4, 1" margins (a letter doesn't need the wide court binding margin).
  * word/numbering.xml multilevel definition; the factual recitals use real numbering.
  * NAMED styles (NoticeHead / AddressBlock / Recital / ChequeLine / SignBlock).
  * Carlito/Caladea fonts (container-safe).
  * docxtpl placeholders keyed to the REAL engine context (client, recipient.name,
    cheques[] loop with amount_words, demand_amount/demand_words, compliance_days),
    each a single run (passes the split-run gate).

Output:  python templates/build_gold_s138_notice.py --out DIR   (default: templates/gold/)
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

A4_W, A4_H = 11906, 16838
MARGIN = 1440  # 1.0" all round (letter, not a bound court filing)
BODY_FONT = "Carlito"
HEADING_FONT = "Caladea"


def _set_a4(section):
    section.page_width = Twips(A4_W)
    section.page_height = Twips(A4_H)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Twips(MARGIN))


def _add_style(doc, name, *, base="Normal", font=BODY_FONT, size=12, bold=False,
               align=None, space_after=6):
    from docx.enum.style import WD_STYLE_TYPE
    styles = doc.styles
    style = styles[name] if name in [s.name for s in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base in [s.name for s in styles]:
        style.base_style = styles[base]
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    if align is not None:
        style.paragraph_format.alignment = align
    style.paragraph_format.space_after = Pt(space_after)
    return style


def _inject_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstractNum = OxmlElement("w:abstractNum")
    abstractNum.set(qn("w:abstractNumId"), "200")
    m = OxmlElement("w:multiLevelType"); m.set(qn("w:val"), "hybridMultilevel"); abstractNum.append(m)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
    pPr = OxmlElement("w:pPr"); ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "360"); pPr.append(ind); lvl.append(pPr)
    abstractNum.append(lvl)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), "20")
    aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), "200"); num.append(aid)
    numbering.insert(0, num); numbering.insert(0, abstractNum)
    return 20


def _numbered(doc, num_id, text):
    p = doc.add_paragraph(style="Recital")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); numPr.append(il)
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id)); numPr.append(ni)
    pPr.append(numPr)
    p.add_run(text)
    return p


def build(out_dir: Path) -> Path:
    doc = Document()
    _set_a4(doc.sections[0])
    _add_style(doc, "NoticeHead", font=HEADING_FONT, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_style(doc, "AddressBlock", size=12, space_after=2)
    _add_style(doc, "Recital", size=12, space_after=6)
    _add_style(doc, "ChequeLine", size=12, space_after=2)
    _add_style(doc, "SignBlock", size=12, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    doc.styles["Normal"].font.name = BODY_FONT
    doc.styles["Normal"].font.size = Pt(12)
    num_id = _inject_numbering(doc)

    # Rule-6 banners.
    doc.add_paragraph("{% if review_banner %}{{ review_banner }}{% endif %}")
    doc.add_paragraph("{% if unverified_marker %}{{ unverified_marker }}{% endif %}")

    # Advocate reference line + date.
    doc.add_paragraph("{{ advocate }}, Advocate", style="AddressBlock")
    doc.add_paragraph("{% if advocate_enrolment %}Enrolment No. {{ advocate_enrolment }}{% endif %}", style="AddressBlock")
    doc.add_paragraph("Date: {{ notice_date }}", style="AddressBlock")
    doc.add_paragraph("", style="AddressBlock")

    # Addressee (the drawer / recipient).
    doc.add_paragraph("To,", style="AddressBlock")
    doc.add_paragraph("{{ recipient.name }},", style="AddressBlock")
    doc.add_paragraph("{{ recipient.address }}", style="AddressBlock")
    doc.add_paragraph("", style="AddressBlock")

    # Subject.
    doc.add_paragraph(
        "Subject: Statutory notice under Section 138 of the Negotiable Instruments "
        "Act, 1881", style="NoticeHead")
    doc.add_paragraph("", style="AddressBlock")
    doc.add_paragraph("Sir/Madam,", style="AddressBlock")
    doc.add_paragraph(
        "Under instructions from and on behalf of my client, {{ client }}, I address "
        "you as follows:", style="Recital")

    # Numbered recitals — REAL numbering.
    _numbered(doc, num_id, "That my client is the payee/holder in due course of the "
                           "cheque(s) described below, issued by you in discharge of "
                           "{{ debt_nature }}.")
    _numbered(doc, num_id, "That the following cheque(s) were drawn by you in favour "
                           "of my client:")

    # Cheque particulars — one PARAGRAPH per cheque. docxtpl's {%p ... %} tags loop at
    # the paragraph level, so each cheque becomes its own line (a plain {% for %} inside
    # one paragraph would run them together — \n in a run is not a paragraph break).
    doc.add_paragraph("{%p for c in cheques %}", style="ChequeLine")
    doc.add_paragraph(
        "Cheque No. {{ c.number }} dated {{ c.date }} for Rs. {{ c.amount }}/- "
        "({{ c.amount_words }} rupees only), drawn on {{ c.bank }}, {{ c.branch }}.",
        style="ChequeLine")
    doc.add_paragraph("{%p endfor %}", style="ChequeLine")

    _numbered(doc, num_id, "That the said cheque(s) were presented for encashment on "
                           "{{ presentation_date }} through {{ presentation_bank }}, "
                           "within the period of their validity.")
    _numbered(doc, num_id, "That the said cheque(s) were returned unpaid vide return "
                           "memo dated {{ memo_date }} for the reason "
                           "“{{ dishonour_reason }}”, which memo my client "
                           "received on {{ memo_received_date }}.")
    _numbered(doc, num_id, "That you have thereby committed the offence punishable "
                           "under Section 138 of the Negotiable Instruments Act, 1881.")

    # The statutory demand.
    doc.add_paragraph(
        "I hereby call upon you to pay to my client the total sum of "
        "Rs. {{ demand_amount }}/- ({{ demand_words }} rupees only) within "
        "{{ compliance_days }} days of receipt of this notice, failing which my client "
        "shall be constrained to initiate prosecution against you under Section 138 of "
        "the Negotiable Instruments Act, 1881, entirely at your risk as to costs and "
        "consequences.", style="Recital")

    # Dispatch + signature.
    doc.add_paragraph("This notice is sent by {{ dispatch_mode }} "
                      "{% if dispatch_tracking_placeholder %}({{ dispatch_tracking_placeholder }}){% endif %}.",
                      style="AddressBlock")
    doc.add_paragraph("", style="AddressBlock")
    doc.add_paragraph("Place: {{ place }}", style="AddressBlock")
    doc.add_paragraph("{{ advocate }}, Advocate", style="SignBlock")
    doc.add_paragraph("(Advocate for {{ client }})", style="SignBlock")

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "s138_notice.docx"
    doc.save(str(out))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build the gold-standard s138_notice.docx.")
    ap.add_argument("--out", type=Path, default=Path(__file__).resolve().parent / "gold")
    args = ap.parse_args(argv)
    out = build(args.out)
    print(f"wrote gold reference: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
